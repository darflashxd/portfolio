"""Generate matching, text-based PDF and editable DOCX. Run: py -3.12 scripts/build-resume.py"""
from pathlib import Path
from html import escape

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
PUBLIC = ROOT / "public"

sections = [
    ("SUMMARY", [
        ("text", "Cybersecurity undergraduate focused on digital forensics and defensive security. Develops forensic challenges, privacy tools, and CTF infrastructure; seeking a cybersecurity internship."),
    ]),
    ("EDUCATION", [
        ("heading", "BINUS University, School of Computer Science | Cyber Security | 2024–2028 (expected)"),
        ("text", "Jakarta, Indonesia | GPA: 3.23 / 4.00"),
    ]),
    ("RELEVANT EXPERIENCE", [
        ("heading", "Member | PETIR Cyber Security | Oct 2026–Present"),
        ("text", "Previously Apprentice | Mar–Oct 2026"),
        ("bullet", "Authored IFEST and BeeCTF forensic challenges in Aug 2026, covering LUKS encryption, ext4 carving, and a simulated malicious AUR package."),
        ("bullet", "Documented Linux CLI and FTK Imager investigation paths; built validation solvers to check intended solutions."),
        ("heading", "Deputy Coordinator, R&D | Cyber Security Community, BINUS | Oct 2025–Present"),
        ("bullet", "Coordinate challenge authors across six competition categories and review contributions for technical accuracy and solvability."),
        ("bullet", "Built GitHub Actions workflows to sync challenge metadata to CTFd via ctfcli and deploy Docker services to a VPS over SSH."),
        ("heading", "Project Officer | National Cyber Week 2025, CSC BINUS | Oct–Dec 2025"),
        ("bullet", "Led competition operations, an on-site IoT hacking final, and workshops on web penetration testing and adversarial AI."),
        ("heading", "R&D Activist | Cyber Security Community, BINUS | Feb–Sep 2025"),
        ("bullet", "Authored file-carving forensics challenges for CSC CTF 2025 and mentored junior members on Volatility and Wireshark investigation workflows."),
    ]),
    ("SELECTED PROJECTS", [
        ("heading", "SafeShare — Privacy tooling for AI conversations"),
        ("link", "github.com/darflashxd/safeshare"),
        ("bullet", "Built a browser extension that detects and redacts sensitive information (credentials, PII, API keys) in AI conversations, with a review interface for adjusting redactions before sharing."),
        ("bullet", "Implemented client-side AES-256-GCM encryption via WebCrypto and a sharing service on Cloudflare Workers with ephemeral burn-after-reading links. Available for users; actively developed."),
        ("heading", "The Poisoned AUR — BeeCTF 2026 digital forensics challenge"),
        ("link", "github.com/PETIRsec/beefest-2026-qual"),
        ("bullet", "Simulated an Arch Linux package compromise requiring correlation across package logs, shell history, systemd journal, ELF binaries, and extended file attributes (xattr) to reconstruct the full attack timeline."),
        ("heading", "PulmoAI — Tuberculosis chest X-ray detection"),
        ("link", "github.com/darflashxd/pulmoai"),
        ("bullet", "Replaced a VGG-style CNN with MobileNetV2 using transfer learning for a chest X-ray classification app with Flask and React."),
    ]),
    ("TECHNICAL SKILLS", [
        ("text", "Forensics & IR: Volatility 3, Autopsy, FTK Imager, Eric Zimmerman Tools, Wireshark; Windows/Linux artifact analysis"),
        ("text", "Detection & SIEM: Splunk, Sysmon, Windows Event Logs (EVTX), YARA Rules, Sigma Rules"),
        ("text", "Development: Python, Bash, TypeScript, Git, Docker, GitHub Actions, ctfcli"),
        ("text", "Reverse Engineering: Ghidra, x64dbg, JADX, APKTool"),
        ("text", "Frameworks: MITRE ATT&CK, NIST SP 800-61, OWASP Top 10"),
    ]),
    ("AWARDS", [
        ("text", "Finalist — Wreck-IT 7.0, General Capture The Flag | Poltek SSN | Aug 2026 | Team HM Plenger"),
        ("link", "darflashxd.my.id/certificates/cert-wreckit-finalist.pdf"),
    ]),
    ("PROFESSIONAL DEVELOPMENT", [
        ("text", "Blue Team Level 1 (BTL1), Security Blue Team — exam preparation in progress; certification not yet earned."),
    ]),
]

name = "AHMAD RAFI SUTANTO"
contact = "Tangerang & Jakarta, Indonesia | +62 895-3831-57017 | rafisutanto@gmail.com"
profiles = "linkedin.com/in/rafisutanto | github.com/darflashxd | darflashxd.my.id"

def add_link(paragraph, label, url):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = label
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)

def link_url(label):
    if '@' in label:
        return 'mailto:' + label
    if label.startswith('+62'):
        return 'tel:' + label.replace(' ', '').replace('-', '')
    if '.' in label:
        return 'https://' + label
    return None

doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(0.5)
section.left_margin = section.right_margin = Inches(0.6)
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Calibri", Pt(10.5)
normal.paragraph_format.space_after = Pt(2)
normal.paragraph_format.line_spacing = 1.0
for text in [name, contact, profiles]:
    paragraph = doc.add_paragraph(text)
    paragraph.alignment = 1
    if text == name:
        paragraph.runs[0].bold = True
        paragraph.runs[0].font.size = Pt(16)
    else:
        paragraph.clear()
        for index, label in enumerate(text.split(' | ')):
            if index:
                paragraph.add_run(' | ')
            url = link_url(label)
            if url:
                add_link(paragraph, label, url)
            else:
                paragraph.add_run(label)

pdfmetrics.registerFont(TTFont("Calibri", "C:/Windows/Fonts/calibri.ttf"))
pdfmetrics.registerFont(TTFont("Calibri-Bold", "C:/Windows/Fonts/calibrib.ttf"))
base = ParagraphStyle("body", fontName="Calibri", fontSize=10.5, leading=12, spaceAfter=2)
styles = {
    "text": base,
    "bullet": ParagraphStyle("bullet", parent=base, leftIndent=10, firstLineIndent=-7),
    "heading": ParagraphStyle("heading", parent=base, fontName="Calibri-Bold", spaceBefore=2, keepWithNext=True),
    "link": ParagraphStyle("link", parent=base, fontSize=10, leading=11, textColor=colors.HexColor("#244970")),
}
story = [Paragraph(name, ParagraphStyle("name", parent=base, fontName="Calibri-Bold", fontSize=16, leading=19, alignment=TA_CENTER))]
for text in [contact, profiles]:
    content = ' | '.join(f'<link href="{escape(link_url(label), quote=True)}">{escape(label)}</link>' if link_url(label) else escape(label) for label in text.split(' | '))
    story.append(Paragraph(content, ParagraphStyle("contact" + text[:3], parent=base, fontSize=10, leading=11, alignment=TA_CENTER)))
story.append(Spacer(1, 3))

for title, entries in sections:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.add_run(title).bold = True
    paragraph.paragraph_format.keep_with_next = True
    story.append(Paragraph(title, ParagraphStyle(title, parent=base, fontName="Calibri-Bold", fontSize=10.5, spaceBefore=5, spaceAfter=2, keepWithNext=True)))
    for kind, text in entries:
        paragraph = doc.add_paragraph(("• " if kind == "bullet" else "") + text)
        if kind == "heading":
            paragraph.runs[0].bold = True
            paragraph.paragraph_format.keep_with_next = True
        content = escape(text)
        if kind == "link":
            content = f'<link href="https://{text}">{content}</link>'
            paragraph.clear()
            add_link(paragraph, text, "https://" + text)
        story.append(Paragraph(("• " if kind == "bullet" else "") + content, styles[kind]))

pdf_path = PUBLIC / "Resume_Ahmad_Rafi_Sutanto.pdf"
SimpleDocTemplate(str(pdf_path), pagesize=(595.28, 841.89), leftMargin=36, rightMargin=36, topMargin=30, bottomMargin=30).build(story)
reader = PdfReader(pdf_path)
assert len(reader.pages) == 1, f"Resume must fit one page, got {len(reader.pages)}"
text = reader.pages[0].extract_text()
assert all(term in text for term in [name, "SafeShare", "not yet earned", "PETIR", "PulmoAI", "MITRE", "YARA", "Ghidra", "OWASP"])
assert "Designed 30+" not in text
assert '\x7f' not in text and '\ufffd' not in text
assert text.count('Finalist') == 1
uris = [a.get_object()['/A']['/URI'] for p in reader.pages for a in p.get('/Annots', [])]
assert len(uris) == 9 and all(u.startswith(('https://', 'mailto:', 'tel:')) for u in uris)
assert any(u.endswith('cert-wreckit-finalist.pdf') for u in uris)
doc.save(PUBLIC / "CV_Ahmad_Rafi_Sutanto.docx")
print("PASS: one-page, text-readable PDF and editable DOCX generated")
